"""
분야(9개 인사 분야) ↔ 과정 연결 검수 시트를 만든다. 엑셀과 구글 시트 모두에서 그대로 쓸 수 있다.

고객용 화면(scripts/build_customer_platform_data.py)은 과정명 키워드로 과정을 9개 분야에 자동 연결한다.
팀원은 과정마다 "가장 맞는 분야"를 드롭다운에서 하나만 고르고, 자동 연결과 맞는지는 수식이 비교한다.

출력: data/hr/output/recommendation/domain_mapping_review_sheet.xlsx
  - 안내    : 입력 방법 5줄과 예시
  - 검수    : 과정 1행. 노란 칸 '가장 맞는 분야' 하나만 입력 (메모는 선택)
  - 교차검수: 20개 과정을 다른 팀원이 한 번 더 고름 (평가자 간 일치도 측정용)
  - 요약    : 진행률, 자동 연결 일치율, 분야별 일치율, 교차검수 일치율·카파 (수식)
  - 분야 기준: 고를 수 있는 분야 목록과 연결 키워드 (드롭다운 목록의 원본)

구글 시트 호환: 시트 이름은 수식에서 따옴표로 감싸고, 드롭다운은 '분야 기준' 시트의 범위를 가리킨다.
구글 드라이브에 올린 뒤 'Google 스프레드시트로 열기'를 하면 드롭다운과 수식이 그대로 동작한다.

이미 입력된 시트를 덮어쓰지 않도록, 파일이 있으면 멈춘다. 다시 만들려면 --force.
실행: python scripts/build_domain_mapping_review_sheet.py [--force]
"""
import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_customer_platform_data import DOMAIN_KEYWORDS, load_courses  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "hr" / "output" / "recommendation" / "domain_mapping_review_sheet.xlsx"

DOMAIN_NAMES = {
    "DOM_01": "인사전략·인력기획", "DOM_02": "조직설계·직무분석", "DOM_03": "인재채용",
    "DOM_04": "온보딩·직원경험", "DOM_05": "성과평가·목표관리", "DOM_06": "보상·급여·복리후생",
    "DOM_07": "인재육성·사내교육", "DOM_08": "조직문화·사내소통", "DOM_09": "노사관계·노무",
}
OFFICE = "사무 일반(엑셀·문서·소통)"
UNRELATED = "인사 업무와 무관"
CHOICES = list(DOMAIN_NAMES.values()) + [OFFICE, UNRELATED]
CHOICE_DESC = {
    OFFICE: "인사 업무에 두루 쓰지만 9개 분야 어디에도 딱 맞지 않는 과정 (엑셀, 보고서, 비즈니스 매너 등)",
    UNRELATED: "인사 담당자의 업무와 관계없는 과정",
}
REVIEWERS = ["팀원A", "팀원B", "팀원C", "팀원D"]
N_CROSS = 20  # 교차검수 과정 수 (팀원 1명당 5개 추가)

FONT = "Arial"
F_BODY = Font(name=FONT, size=10)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True)
F_LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
FILL_HEAD = PatternFill("solid", fgColor="1F4E78")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_AUTO = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

# '분야 기준' 시트의 선택지 범위 (드롭다운과 카파 계산에 사용)
CHOICE_RANGE = f"'분야 기준'!$B$2:$B${len(CHOICES) + 1}"


def q(sheet):
    """구글 시트 호환을 위해 시트 이름을 항상 따옴표로 감싼다."""
    return f"'{sheet}'"


def header(ws, columns):
    for i, (name, width, _) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30


def match_formula(pick, auto):
    """고른 분야가 자동 연결 분야와 맞는지. 자동 연결이 '(없음)'이면 사무 일반·무관을 고른 경우가 일치."""
    return (f'=IF({pick}="","",IF({auto}="(없음)",'
            f'IF(OR({pick}="{OFFICE}",{pick}="{UNRELATED}"),"일치","불일치"),'
            f'IF(ISNUMBER(SEARCH({pick},{auto})),"일치","불일치")))')


def add_choice_dropdown(ws, col_letter, last):
    dv = DataValidation(type="list", formula1=f"={CHOICE_RANGE}", allow_blank=True,
                        showErrorMessage=True, errorTitle="목록에서 고르세요",
                        error="'분야 기준' 시트의 분야 중 하나를 고르세요")
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{last}")


# ------------------------------------------------------------------------------
# 검수 시트
# ------------------------------------------------------------------------------
REVIEW_COLUMNS = [
    ("번호", 6, False), ("배정", 8, False), ("과정명", 46, False), ("기관", 20, False), ("링크", 8, False),
    ("자동 연결 분야", 26, False), ("가장 맞는 분야 (고르기)", 26, True), ("메모 (선택)", 30, True),
    ("자동 연결과 비교", 12, False), ("과정ID", 26, False),
]
RCOL = {name: i + 1 for i, (name, _, _) in enumerate(REVIEW_COLUMNS)}


def rl(name):
    return get_column_letter(RCOL[name])


def auto_text(c):
    return ", ".join(DOMAIN_NAMES[d] for d in c["doms"]) or "(없음)"


def build_review_sheet(ws, courses):
    ws.title = "검수"
    header(ws, REVIEW_COLUMNS)
    for r, c in enumerate(courses, start=2):
        values = {
            "번호": r - 1, "배정": REVIEWERS[(r - 2) % len(REVIEWERS)], "과정명": c["title"], "기관": c["inst"],
            "링크": "열기" if c["url"] else "", "자동 연결 분야": auto_text(c),
            "자동 연결과 비교": match_formula(f"{rl('가장 맞는 분야 (고르기)')}{r}", f"{rl('자동 연결 분야')}{r}"),
            "과정ID": c["id"],
        }
        for name, (_, _, is_input) in zip(RCOL, REVIEW_COLUMNS):
            cell = ws.cell(row=r, column=RCOL[name], value=values.get(name))
            cell.font, cell.border, cell.alignment = F_BODY, BORDER, WRAP
            cell.fill = FILL_INPUT if is_input else FILL_AUTO
        if c["url"]:
            link = ws.cell(row=r, column=RCOL["링크"])
            link.hyperlink, link.font = c["url"], F_LINK
    last = len(courses) + 1
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(REVIEW_COLUMNS))}{last}"
    add_choice_dropdown(ws, rl("가장 맞는 분야 (고르기)"), last)
    cmp_col = rl("자동 연결과 비교")
    ws.conditional_formatting.add(f"{cmp_col}2:{cmp_col}{last}",
                                  FormulaRule(formula=[f'{cmp_col}2="불일치"'], font=Font(name=FONT, color="C00000", bold=True)))
    return last


# ------------------------------------------------------------------------------
# 교차검수 시트
# ------------------------------------------------------------------------------
CROSS_COLUMNS = [
    ("번호", 6, False), ("교차 배정", 9, False), ("과정명", 46, False), ("기관", 20, False), ("링크", 8, False),
    ("자동 연결 분야", 26, False), ("가장 맞는 분야 (고르기)", 26, True), ("메모 (선택)", 30, True),
    ("과정ID", 26, False), ("1차 배정", 9, False), ("1차 선택", 22, False), ("두 사람 일치", 10, False),
]
CCOL = {name: i + 1 for i, (name, _, _) in enumerate(CROSS_COLUMNS)}


def cl(name):
    return get_column_letter(CCOL[name])


def pick_cross_rows(courses, n=N_CROSS):
    """분야(첫 번째 연결 분야, 없으면 '(없음)')별로 돌아가며 고르게 n개 과정을 고른다. 결과는 항상 같다."""
    groups = {}
    for i, c in enumerate(courses):
        groups.setdefault(c["doms"][0] if c["doms"] else "NONE", []).append(i)
    order = list(DOMAIN_KEYWORDS) + ["NONE"]
    queues = {g: groups[g][:: max(1, len(groups[g]) // 4)] + groups[g] for g in order if g in groups}
    picked, seen = [], set()
    while len(picked) < n and any(queues.values()):
        for g in order:
            qq = queues.get(g, [])
            while qq and qq[0] in seen:
                qq.pop(0)
            if qq and len(picked) < n:
                seen.add(qq[0])
                picked.append(qq.pop(0))
    return sorted(picked)


def assign_cross(rows):
    """1차 배정과 다른 팀원에게, 팀원별 개수가 고르도록 배정한다."""
    load = {m: 0 for m in REVIEWERS}
    second_of = {}
    for idx in rows:
        first = REVIEWERS[idx % len(REVIEWERS)]
        second = min((m for m in REVIEWERS if m != first), key=lambda m: (load[m], REVIEWERS.index(m)))
        load[second] += 1
        second_of[idx] = second
    target = -(-len(rows) // len(REVIEWERS))
    moved = True
    while moved:
        moved = False
        over = [m for m in REVIEWERS if load[m] > target or (load[m] == target and min(load.values()) < target - 1)]
        under = [m for m in REVIEWERS if load[m] < max(load.values()) - 1]
        for idx in rows:
            src = second_of[idx]
            dst = next((m for m in under if m != REVIEWERS[idx % len(REVIEWERS)]), None)
            if src in over and dst and load[src] - load[dst] > 1:
                second_of[idx], load[src], load[dst], moved = dst, load[src] - 1, load[dst] + 1, True
                break
    return second_of


def build_cross_sheet(ws, courses, review_last):
    """1차 담당과 다른 팀원이 같은 과정을 다시 고른다. 1차 선택이 보이는 열은 숨긴다(따라 하기 방지)."""
    ws.title = "교차검수"
    header(ws, CROSS_COLUMNS)
    pick_rng = f"{q('검수')}!${rl('가장 맞는 분야 (고르기)')}$2:${rl('가장 맞는 분야 (고르기)')}${review_last}"
    id_rng = f"{q('검수')}!${rl('과정ID')}$2:${rl('과정ID')}${review_last}"
    rows = pick_cross_rows(courses)
    second_of = assign_cross(rows)
    for r, idx in enumerate(rows, start=2):
        c = courses[idx]
        id_ref = f"${cl('과정ID')}{r}"
        first_pick = f"INDEX({pick_rng},MATCH({id_ref},{id_rng},0))"
        values = {
            "번호": r - 1, "교차 배정": second_of[idx], "과정명": c["title"], "기관": c["inst"],
            "링크": "열기" if c["url"] else "", "자동 연결 분야": auto_text(c), "과정ID": c["id"],
            "1차 배정": REVIEWERS[idx % len(REVIEWERS)],
            "1차 선택": f'=IFERROR(IF({first_pick}="","",{first_pick}),"")',
            "두 사람 일치": (f'=IF(OR({cl("1차 선택")}{r}="",{cl("가장 맞는 분야 (고르기)")}{r}=""),"",'
                         f'IF({cl("1차 선택")}{r}={cl("가장 맞는 분야 (고르기)")}{r},"일치","불일치"))'),
        }
        for name, (_, _, is_input) in zip(CCOL, CROSS_COLUMNS):
            cell = ws.cell(row=r, column=CCOL[name], value=values.get(name))
            cell.font, cell.border, cell.alignment = F_BODY, BORDER, WRAP
            cell.fill = FILL_INPUT if is_input else FILL_AUTO
        if c["url"]:
            link = ws.cell(row=r, column=CCOL["링크"])
            link.hyperlink, link.font = c["url"], F_LINK
    last = len(rows) + 1
    for name in ["과정ID", "1차 배정", "1차 선택", "두 사람 일치"]:
        ws.column_dimensions[cl(name)].hidden = True
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{cl('메모 (선택)')}{last}"
    add_choice_dropdown(ws, cl("가장 맞는 분야 (고르기)"), last)
    return last


# ------------------------------------------------------------------------------
# 안내 / 분야 기준 / 요약
# ------------------------------------------------------------------------------
def build_guide(ws):
    ws.title = "안내"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 95
    rows = [
        ("분야–과정 검수 시트", None),
        (None, None),
        ("할 일", "과정마다 '가장 맞는 분야'를 드롭다운에서 하나 고르면 끝입니다. 자동 연결과 맞는지는 시트가 계산합니다."),
        ("1", "'검수' 시트에서 '배정' 열 필터로 내 이름(팀원A~D)만 봅니다."),
        ("2", "과정명을 읽고 노란 칸 '가장 맞는 분야'에서 하나를 고릅니다. 분야가 여러 개 맞으면 가장 중심이 되는 하나만."),
        ("3", "9개 분야에 안 맞지만 업무에 도움 되면 '사무 일반(엑셀·문서·소통)', 업무와 관계없으면 '인사 업무와 무관'."),
        ("4", "과정명만으로 모르겠으면 '링크'를 열어 소개를 봅니다. 그래도 애매하면 고르고 메모에 이유를 적습니다."),
        ("5", "'교차검수' 시트에서 '교차 배정'이 내 이름인 5개도 같은 방법으로 고릅니다. 다른 사람의 선택은 보지 마세요."),
        (None, None),
        ("주의", "회색 칸은 고치지 마세요. 고용24 과정은 2025년 회차라 링크의 회차가 끝났을 수 있으니 내용만 보고 판단합니다."),
        ("결과", "'요약' 시트에서 진행률, 자동 연결 일치율, 분야별 일치율, 두 사람 일치도가 자동으로 계산됩니다."),
        (None, None),
        ("입력 예시", None),
    ]
    for r, (a, b) in enumerate(rows, start=1):
        ca, cb = ws.cell(row=r, column=1, value=a), ws.cell(row=r, column=2, value=b)
        ca.font, cb.font = F_BOLD, F_BODY
        ca.alignment = cb.alignment = WRAP
    ws["A1"].font = F_TITLE
    examples = [
        ("과정명", "자동 연결 분야", "가장 맞는 분야 (고르기)", "자동 연결과 비교"),
        ("[여의도] 채용면접위원 양성", "인재채용", "인재채용", "일치"),
        ("엑셀을 활용한 보고서 작성", "(없음)", OFFICE, "일치"),
        ("채용에서 퇴직까지 인사노무관리 핵심실무", "인재채용, 노사관계·노무", "노사관계·노무", "일치"),
        ("[DX&AI] 채용 분석 실무", "인재채용", "인사전략·인력기획", "불일치 (메모: 채용 데이터 분석 중심)"),
    ]
    start = len(rows) + 1
    widths = {"C": 26, "D": 34}
    for k, w in widths.items():
        ws.column_dimensions[k].width = w
    for i, row in enumerate(examples):
        for j, v in enumerate(row):
            cell = ws.cell(row=start + i, column=j + 1, value=v)
            cell.border = BORDER
            cell.alignment = WRAP
            if i == 0:
                cell.font, cell.fill = F_HEAD, FILL_HEAD
            else:
                cell.font = F_BODY
                cell.fill = FILL_INPUT if j == 2 else FILL_AUTO
    ws.column_dimensions["A"].width = 34


def build_domains(ws):
    ws.title = "분야 기준"
    heads = [("번호", 6), ("고를 수 있는 분야", 26), ("설명 / 자동 연결 키워드(과정명에 포함되면 연결)", 90)]
    for i, (h, w) in enumerate(heads, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    kw_by_name = {DOMAIN_NAMES[d]: ", ".join(k) for d, k in DOMAIN_KEYWORDS.items()}
    for r, name in enumerate(CHOICES, start=2):
        for i, v in enumerate([r - 1, name, CHOICE_DESC.get(name, "키워드: " + kw_by_name.get(name, ""))], start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font, c.border, c.alignment = F_BODY, BORDER, WRAP


def build_summary(ws, last, cross_last):
    ws.title = "요약"
    ws.column_dimensions["A"].width = 30
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 13
    S = q("검수")
    X = q("교차검수")
    pick = f"{S}!${rl('가장 맞는 분야 (고르기)')}$2:${rl('가장 맞는 분야 (고르기)')}${last}"
    cmp_ = f"{S}!${rl('자동 연결과 비교')}$2:${rl('자동 연결과 비교')}${last}"
    auto = f"{S}!${rl('자동 연결 분야')}$2:${rl('자동 연결 분야')}${last}"
    asg = f"{S}!${rl('배정')}$2:${rl('배정')}${last}"
    xasg = f"{X}!${cl('교차 배정')}$2:${cl('교차 배정')}${cross_last}"
    xpick = f"{X}!${cl('가장 맞는 분야 (고르기)')}$2:${cl('가장 맞는 분야 (고르기)')}${cross_last}"
    xfirst = f"{X}!${cl('1차 선택')}$2:${cl('1차 선택')}${cross_last}"

    def put(cell, value, font=F_BODY, fmt=None):
        ws[cell] = value
        ws[cell].font = font
        ws[cell].border = BORDER
        if fmt:
            ws[cell].number_format = fmt

    ws["A1"] = "검수 요약 (자동 계산)"
    ws["A1"].font = F_TITLE
    put("A3", "전체 과정 수", F_BOLD); put("B3", f"=COUNTA({S}!$C$2:$C${last})")
    put("A4", "입력 완료", F_BOLD); put("B4", f'=COUNTIF({pick},"?*")')
    put("A5", "진행률", F_BOLD); put("B5", "=IF(B3=0,0,B4/B3)", fmt="0.0%")
    put("A6", "자동 연결 일치", F_BOLD); put("B6", f'=COUNTIF({cmp_},"일치")')
    put("A7", "자동 연결 일치율", F_BOLD); put("B7", "=IF(B4=0,0,B6/B4)", fmt="0.0%")

    put("A9", "고른 분야", F_BOLD); put("B9", "건수", F_BOLD); put("C9", "비율", F_BOLD)
    for i, name in enumerate(CHOICES):
        r = 10 + i
        put(f"A{r}", name); put(f"B{r}", f"=COUNTIF({pick},A{r})"); put(f"C{r}", f"=IF($B$4=0,0,B{r}/$B$4)", fmt="0.0%")

    base = 10 + len(CHOICES) + 1
    put(f"A{base}", "자동 연결 분야별", F_BOLD)
    for i, h in enumerate(["자동 연결 과정", "입력 완료", "일치", "일치율"], start=2):
        put(f"{get_column_letter(i)}{base}", h, F_BOLD)
    for i, name in enumerate(list(DOMAIN_NAMES.values()) + ["(없음)"]):
        r = base + 1 + i
        crit = f'"*"&A{r}&"*"' if name != "(없음)" else '"(없음)"'
        put(f"A{r}", name)
        put(f"B{r}", f"=COUNTIF({auto},{crit})")
        put(f"C{r}", f'=COUNTIFS({auto},{crit},{pick},"?*")')
        put(f"D{r}", f'=COUNTIFS({auto},{crit},{cmp_},"일치")')
        put(f"E{r}", f"=IF(C{r}=0,0,D{r}/C{r})", fmt="0.0%")
    r_end = base + 1 + len(DOMAIN_NAMES)
    note = ws.cell(row=r_end + 1, column=1,
                   value="일치율 = 입력한 과정 중 고른 분야가 자동 연결 분야에 들어 있는 비율. '(없음)'은 사무 일반·무관을 고른 경우가 일치.")
    note.font = F_NOTE

    tb = r_end + 3
    put(f"A{tb}", "팀원별", F_BOLD)
    for i, h in enumerate(["배정", "완료", "진행률", "교차 배정", "교차 완료"], start=2):
        put(f"{get_column_letter(i)}{tb}", h, F_BOLD)
    for i, m in enumerate(REVIEWERS):
        r = tb + 1 + i
        put(f"A{r}", m)
        put(f"B{r}", f"=COUNTIF({asg},A{r})")
        put(f"C{r}", f'=COUNTIFS({asg},A{r},{pick},"?*")')
        put(f"D{r}", f"=IF(B{r}=0,0,C{r}/B{r})", fmt="0.0%")
        put(f"E{r}", f"=COUNTIF({xasg},A{r})")
        put(f"F{r}", f'=COUNTIFS({xasg},A{r},{xpick},"?*")')

    # 교차검수: 두 사람 모두 고른 과정만으로 일치율과 카파 계수
    both = f'(({xfirst}<>"")*({xpick}<>""))'
    cb = tb + 1 + len(REVIEWERS) + 1
    put(f"A{cb}", "교차검수 (두 사람 일치도)", F_BOLD)
    put(f"A{cb + 1}", "교차검수 과정 수"); put(f"B{cb + 1}", f"=COUNTA({X}!$C$2:$C${cross_last})")
    put(f"A{cb + 2}", "두 사람 모두 입력"); put(f"B{cb + 2}", f"=SUMPRODUCT({both})")
    put(f"A{cb + 3}", "같은 분야를 고름"); put(f"B{cb + 3}", f"=SUMPRODUCT({both}*({xfirst}={xpick}))")
    put(f"A{cb + 4}", "일치율"); put(f"B{cb + 4}", f"=IF(B{cb + 2}=0,0,B{cb + 3}/B{cb + 2})", fmt="0.0%")
    pe = "+".join(f'SUMPRODUCT({both}*({xfirst}="{v}"))*SUMPRODUCT({both}*({xpick}="{v}"))' for v in CHOICES)
    put(f"A{cb + 5}", "우연 일치 기대값(pe)"); put(f"B{cb + 5}", f"=IF(B{cb + 2}=0,0,({pe})/B{cb + 2}^2)", fmt="0.000")
    put(f"A{cb + 6}", "카파 계수(κ)"); put(f"B{cb + 6}", f'=IF(OR(B{cb + 2}=0,B{cb + 5}=1),"-",(B{cb + 4}-B{cb + 5})/(1-B{cb + 5}))', fmt="0.00")
    n2 = ws.cell(row=cb + 7, column=1, value="카파 0.6 이상이면 팀원 간 기준이 대체로 같음, 0.4 미만이면 기준을 다시 맞추는 회의가 필요합니다.")
    n2.font = F_NOTE


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true", help="이미 있는 시트를 덮어쓰기 (입력 내용이 사라짐)")
    args = ap.parse_args()
    if OUT.exists() and not args.force:
        sys.exit(f"[중단] {OUT.relative_to(ROOT)}가 이미 있습니다. 입력 내용을 지우고 다시 만들려면 --force를 붙이세요.")

    courses = load_courses()
    wb = Workbook()
    build_guide(wb.active)
    last = build_review_sheet(wb.create_sheet(), courses)
    cross_last = build_cross_sheet(wb.create_sheet(), courses, last)
    build_summary(wb.create_sheet(), last, cross_last)
    build_domains(wb.create_sheet())
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"[저장] {OUT.relative_to(ROOT)}: 과정 {len(courses)}개(팀원별 약 {len(courses) // len(REVIEWERS)}개) "
          f"+ 교차검수 {cross_last - 1}개(팀원별 {(cross_last - 1) // len(REVIEWERS)}개)")


if __name__ == "__main__":
    main()
